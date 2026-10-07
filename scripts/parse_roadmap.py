# -*- coding: utf-8 -*-
"""parse_roadmap.py — дорожные карты развития из xlsx в scripts/roadmaps.json.

ЧТО ЧИТАЕТ
    Папка G:\LMStudio\Hermes\hermes-workspace\itogi\Дорожные карты развития\
    Файл «Дорожная карта Roadmap - <Имя Фамилия>.xlsx», лист выбранного варианта
    («Выбор сотрудника» — приоритет, далее «Целевая», «Альтернативная»).

СТРУКТУРА ЛИСТА (как в файле Болгова)
    Строка 1  — «🗺 Карта №N: «Название»»
    Строка 3  — «🎯 Глобальная цель: …»
    Строки 5+ — «💎 Ценность для тебя:» и пункты «· …»
    Строка 10 — заголовки: Этап | Подцели и Ценность | Теория+Практика | Инструкция
    Далее блоки этапов: 1-я строка — «N. Название», 2-я — «(подзаголовок)»,
    затем строки «1. подцель», «Теория: …», «Практика: …», «Ценность: …»,
    «Где брать: …» / «Как делать: …», «🧘: …» (мягкий навык / самопроверка).

ЗАПУСК
    python parse_roadmap.py          # разобрать папку и записать scripts/roadmaps.json
    python parse_roadmap.py --show   # только показать результат
"""
import io
import json
import os
import re
import sys

import openpyxl

HERMES = r"G:\LMStudio\Hermes"
ROAD_DIR = os.path.join(HERMES, "hermes-workspace", "itogi", "Дорожные карты развития")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "roadmaps.json")
STAFF = os.path.join(HERE, "staff.json")
SHEET_PRIORITY = ["Выбор сотрудника", "Целевая", "Альтернативная"]
SHEET_LABEL = {"Выбор сотрудника": "выбранный сотрудником вариант",
               "Целевая": "рекомендация руководителя",
               "Альтернативная": "альтернативный трек"}


def staff_map():
    """{фамилия: slug} — из staff.json (снимок сотрудников сайта)."""
    try:
        items = json.load(io.open(STAFF, encoding="utf-8"))
    except Exception:
        return {}
    out = {}
    for e in items:
        full, slug = (e.get("fullName") or ""), (e.get("slug") or "")
        if full and slug:
            out[full.split(" ")[0].lower()] = slug
    return out


def cells(ws, row):
    """Значения строки по колонкам A..E (пустое -> '')."""
    out = []
    for c in range(1, 6):
        v = ws.cell(row=row, column=c).value
        out.append(str(v).strip() if v is not None and str(v).strip() else "")
    return out


def clean(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def parse_sheet(ws):
    head = cells(ws, 1)[0]
    m = re.search(r"Карта\s*№?\s*(\d+)\s*[:.]?\s*", head)
    number = int(m.group(1)) if m else None
    rest = head[m.end():] if m else head
    rest = rest.strip().strip('«»"').strip()
    title, subtitle = rest, ""
    if ":" in rest:
        a, b = rest.split(":", 1)
        if len(b.strip()) > 3:
            title, subtitle = a.strip(), b.strip()

    goal = re.sub(r"^[^\w]*Глобальная цель\s*:\s*", "", cells(ws, 3)[0]).strip()
    values = []
    for r in range(5, min(ws.max_row, 12) + 1):
        line = cells(ws, r)[0]
        if line.startswith("·"):
            values.append(line.lstrip("·").strip())
        elif values and line and not line.startswith("💎"):
            break

    stages = []
    cur = None
    for r in range(11, ws.max_row + 1):
        a, b, c, d, e = cells(ws, r)
        if not any([b, c, d, e]):
            continue
        if b.startswith("(") and cur is not None:          # подзаголовок этапа
            cur["tag"] = clean(b.strip("()"))              # (остальные колонки строки ещё читаем)
        mm = re.match(r"^(\d+)[.)]\s*(.+)$", b)
        if mm:                                             # новый этап
            cur = {"n": int(mm.group(1)), "title": clean(mm.group(2)), "tag": "",
                   "goals": [], "value": "", "theory": "", "practice": "", "how": "", "mind": ""}
            stages.append(cur)
        if cur is None:
            continue
        for col in (c, d, e):
            if not col:
                continue
            low = col.lower()
            if low.startswith("ценность"):
                cur["value"] = clean(cur["value"] + " " + re.sub(r"^Ценность\s*:?\s*", "", col))
            elif col.startswith("🧘"):
                body = re.sub(r"^🧘[^A-Za-zА-Яа-я0-9]*\s*", "", col).strip()
                cur["mind"] = clean(cur["mind"] + " " + body)
            elif low.startswith("теория"):
                cur["theory"] = clean(cur["theory"] + " " + re.sub(r"^Теория\s*:?\s*", "", col))
            elif low.startswith("практика"):
                cur["practice"] = clean(cur["practice"] + " " + re.sub(r"^Практика\s*:?\s*", "", col))
            elif low.startswith("где брать") or low.startswith("как делать"):
                cur["how"] = clean(cur["how"] + " " + col)
            elif re.match(r"^\d+[.)]\s*\S", col):
                g = clean(re.sub(r"^\d+[.)]\s*", "", col))
                if g and g not in cur["goals"]:
                    cur["goals"].append(g)
    for i, s in enumerate(stages, 1):
        s["rawN"] = s["n"]
        s["n"] = i                                   # нумерация строго по порядку (в файле бывают повторы)
        for k in ("value", "theory", "practice", "how", "mind"):
            s[k] = clean(s[k])
    return {"number": number, "title": title, "subtitle": subtitle,
            "goal": goal, "values": values, "stages": stages}


def parse_file(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    variants = {}
    for name in wb.sheetnames:
        if name in SHEET_PRIORITY:
            variants[name] = parse_sheet(wb[name])
    if not variants:
        return None
    for name in SHEET_PRIORITY:                     # приоритет выбранному варианту
        if name in variants:
            card = dict(variants[name])
            card["sheet"] = name
            card["sheetLabel"] = SHEET_LABEL.get(name, name)
            card["variants"] = [{"sheet": k, "title": v.get("title"), "stages": len(v.get("stages", []))}
                                for k, v in variants.items()]
            return card
    return None


def main(argv):
    smap = staff_map()
    result = {}
    if not os.path.isdir(ROAD_DIR):
        print("нет папки:", ROAD_DIR)
        return 1
    for fn in sorted(os.listdir(ROAD_DIR)):
        if not fn.lower().endswith((".xlsx", ".xlsm")) or fn.startswith("~$"):
            continue
        slug, surname = "", ""
        tail = fn.rsplit("-", 1)[-1].rsplit(".", 1)[0]
        for word in re.findall(r"[А-ЯЁA-Z][а-яёa-z]+", tail):
            if word.lower() in smap:
                slug, surname = smap[word.lower()], word.lower()
                break
        card = parse_file(os.path.join(ROAD_DIR, fn))
        if not card:
            continue
        card["source"] = fn
        card["name"] = re.sub(r"\s+", " ", tail).strip()
        key = slug or ("file:" + fn)
        result[key] = card
        print("%s: %s (slug %s) - этапов %d, карта N%s, лист %s"
              % (fn, surname or "?", slug or "НЕ НАЙДЕН", len(card["stages"]), card["number"], card["sheet"]))
    if "--show" not in argv:
        io.open(OUT, "w", encoding="utf-8").write(json.dumps(result, ensure_ascii=False, indent=1))
        print("записано:", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
